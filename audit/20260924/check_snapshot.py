"""Read-only checks of the frozen submission; writes evidence only in this folder.

AI-assisted audit: OpenAI Codex. Original solver and results are not modified.
Run from the repository root: python audit/20260924/check_snapshot.py
"""
from pathlib import Path
from collections import Counter
import hashlib
import importlib.metadata
import json
import sys

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
import solve_d


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    manifest = json.loads((ROOT / 'results/复现清单.json').read_text(encoding='utf-8'))
    evidence = {'scope': 'Frozen files and arithmetic/resource checks, not a full communication proof or optimization rerun',
                'python': sys.version, 'dependencies': {n: importlib.metadata.version(n) for n in
                    ('numpy', 'pandas', 'Pillow', 'openpyxl', 'matplotlib')}, 'manifest': {}}
    for category in ('input_files', 'source_files', 'output_files'):
        records = []
        for item in manifest[category]:
            path = ROOT / item['path'].replace('\\', '/')
            actual = hashlib.sha256(path.read_bytes()).hexdigest().upper() if path.exists() else None
            lf_match = category == 'source_files' and path.exists() and hashlib.sha256(path.read_bytes().replace(b'\r\n', b'\n')).hexdigest().upper() == item['sha256']
            records.append({'path': item['path'], 'match': actual == item['sha256'], 'matches_after_lf_normalization': lf_match,
                            'expected': item['sha256'], 'actual': actual})
        evidence['manifest'][category] = {'count': len(records), 'matched': sum(r['match'] for r in records),
                                         'mismatches': [r for r in records if not r['match']]}
    data = solve_d.load_inputs(ROOT)
    evidence['inputs'] = {'nodes': len(data['nodes']), 'boxes': len(data['boxes']),
                          'summary_rows': len(data['summary'])}
    box_ids = set(data['boxes']['货箱编号'])
    checks = {}
    csv = lambda name: pd.read_csv(ROOT / 'results' / name)
    for q in ('q1', 'q2', 'q3'):
        trips = csv('q1_batches.csv' if q == 'q1' else q + '_transport_trips.csv')
        coverage = Counter(b for ids in trips['box_ids'] for b in ids.split(';'))
        checks[q] = {'trips': len(trips), 'box_coverage_exact': set(coverage) == box_ids and all(n == 1 for n in coverage.values()),
                     'energy_kwh': float(trips.energy_kwh.sum()), 'minimum_return_soc': float(trips.return_soc.min())}
        errors = []
        for row in trips.to_dict('records'):
            boxes = data['boxes'][data['boxes']['货箱编号'].isin(row['box_ids'].split(';'))]
            g = data['transport_types'][row['type_id']]
            # Uses the existing route physics, but rechecks stored CSV values independently of JSON sanity flags.
            if q == 'q1':
                calc = solve_d.batch_metrics(data['dem'], data['nodes']['O01'], data['nodes'][row['service_id']], g, boxes)
                error = abs(calc['energy_kwh'] - row['energy_kwh'])
                if error > 1e-8: errors.append({'trip': row['trip_id'], 'energy_error': error})
            if boxes['单箱质量（kg）'].sum() > g.payload_kg + 1e-8:
                errors.append({'trip': row['trip_id'], 'error': 'mass exceeds capacity'})
            if boxes['单箱体积（m³）'].sum() > g.volume_m3 + 1e-8:
                errors.append({'trip': row['trip_id'], 'error': 'volume exceeds capacity'})
            if row['return_soc'] < g.reserve - 1e-8:
                errors.append({'trip': row['trip_id'], 'error': 'return SOC below reserve'})
        checks[q]['errors'] = errors
        if q != 'q1':
            deliveries = csv(q + '_box_deliveries.csv')
            hard = deliveries[deliveries.hard_deadline_s.notna()]
            checks[q].update({'delivery_coverage_exact': len(deliveries) == len(box_ids) and set(deliveries.box_id) == box_ids,
                              'hard_boxes': len(hard), 'hard_late_recomputed': int((hard.completion_s > hard.hard_deadline_s + 1e-8).sum()),
                              'minimum_hard_slack_s': float((hard.hard_deadline_s - hard.completion_s).min()),
                              'return_max_s': float(trips.return_s.max())})
            conflicts = []
            for key, end in [('unit_id', 'return_s'), ('battery_id', 'battery_charge_complete_s')]:
                for entity, group in trips.groupby(key):
                    previous_end = -float('inf')
                    for row in group.sort_values('start_s').to_dict('records'):
                        if row['start_s'] < previous_end - 1e-8: conflicts.append([key, entity, row['trip_id']])
                        previous_end = max(previous_end, row[end])
            checks[q]['resource_conflicts'] = conflicts
    relay = csv('q3_relay_missions.csv')
    conflicts = []
    for key, end in [('relay_unit_id', 'available_s'), ('energy_pack_id', 'pack_charge_complete_s')]:
        for entity, group in relay.groupby(key):
            previous_end = -float('inf')
            for row in group.sort_values('start_s').to_dict('records'):
                if row['start_s'] < previous_end - 1e-8: conflicts.append([key, entity, row['mission_id']])
                previous_end = max(previous_end, row[end])
    checks['relay'] = {'missions': len(relay), 'resource_conflicts': conflicts,
                       'units_used': relay.relay_unit_id.nunique(), 'packs_used': relay.energy_pack_id.nunique(),
                       'minimum_return_soc': float(relay.return_soc.min()), 'return_max_s': float(relay.return_s.max())}
    gap = csv('q4_shortage_redundancy.csv')
    checks['q4'] = {'shortfall_formula': bool((gap.shortfall == (gap.total_group_need - gap.inventory).clip(lower=0)).all()),
                    'redundancy_formula': bool((gap.partition_redundancy == gap.total_group_need - gap.central_need).all())}
    evidence['stored_result_checks'] = checks
    wb = load_workbook(ROOT / '结果提交表_D题.xlsx', read_only=True, data_only=True)
    evidence['workbook'] = {ws.title: {'nonempty_data_rows': sum(any(v is not None for v in row) for row in ws.iter_rows(min_row=2, values_only=True)),
                                      'headers': list(next(ws.iter_rows(max_row=1, values_only=True)))} for ws in wb}
    wb.close()
    evidence['flowcharts'] = [p.name for p in (ROOT / 'figures').glob('flow_*.png')]
    evidence['paper_at_root'] = [p.name for p in ROOT.glob('*.docx')]
    (OUT / 'snapshot_checks.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(evidence, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
