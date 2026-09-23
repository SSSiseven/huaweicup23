"""Fill the official result workbook from frozen CSV outputs.

AI 辅助说明：本程序及代码是在人工智能工具 OpenAI Codex（GPT-5，
开发机构 OpenAI，版本发布日期 2025-08-07）辅助下完成的。
所有字段映射和提交内容均须由参赛者独立复核后使用。
"""

from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "D题" / "结果提交模板.xlsx"
OUTPUT = ROOT / "结果提交表_D题.xlsx"
RESULTS = ROOT / "results"


def rows_from(df: pd.DataFrame, columns: list[str], transforms=None):
    transforms = transforms or {}
    for _, row in df.iterrows():
        values = []
        for col in columns:
            value = row[col]
            if pd.isna(value):
                value = None
            if col in transforms:
                value = transforms[col](value, row)
            if hasattr(value, "item"):
                value = value.item()
            values.append(value)
        yield values


def fill_sheet(ws, rows: list[list], ncols: int) -> None:
    # The official template contains blank formatted rows. Preserve row 1 and
    # use row 2 as the formatting prototype for all generated records.
    prototype = [ws.cell(2, c) for c in range(1, ncols + 1)]
    for r in range(2, max(ws.max_row, len(rows) + 1) + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None
    for r_idx, values in enumerate(rows, 2):
        for c_idx, value in enumerate(values, 1):
            cell = ws.cell(r_idx, c_idx, value)
            src = prototype[c_idx - 1]
            if src.has_style:
                cell._style = copy(src._style)
            cell.number_format = src.number_format
            cell.alignment = copy(src.alignment) if src.alignment else Alignment()
    ws.auto_filter.ref = f"A1:{ws.cell(1, ncols).column_letter}{len(rows)+1}"
    ws.freeze_panes = "A2"


def main() -> None:
    wb = load_workbook(TEMPLATE)

    q1 = pd.read_csv(RESULTS / "q1_batches.csv")
    fill_sheet(wb["Q1_单点组批"], list(rows_from(q1,
        ["trip_id", "service_id", "type_id", "box_ids", "mass_kg", "volume_m3", "duration_s", "energy_kwh", "return_soc"],
        {"return_soc": lambda v, _: 100 * float(v)})), 9)

    q2t = pd.read_csv(RESULTS / "q2_transport_trips.csv")
    fill_sheet(wb["Q2_运输架次"], list(rows_from(q2t,
        ["trip_id", "unit_id", "type_id", "battery_id", "start_s", "route", "return_s", "energy_kwh"])), 8)

    q2d = pd.read_csv(RESULTS / "q2_box_deliveries.csv")
    fill_sheet(wb["Q2_逐箱交付"], list(rows_from(q2d,
        ["box_id", "trip_id", "service_id", "completion_s"])), 4)

    q3r = pd.read_csv(RESULTS / "q3_relay_missions.csv")
    fill_sheet(wb["Q3_中继架次"], list(rows_from(q3r,
        ["mission_id", "relay_unit_id", "energy_pack_id", "start_s", "lon", "lat", "altitude_m",
         "link_ready_s", "service_end_s", "return_s", "energy_kwh"])), 11)

    q3c = pd.read_csv(RESULTS / "q3_communication_segments.csv")
    q3c["assurance_clean"] = q3c["phase"].astype(str).map(lambda x: "直连" if x == "direct" else "中继")
    fill_sheet(wb["Q3_通信保障"], list(rows_from(q3c,
        ["trip_id", "phase", "start_s", "end_s", "assurance_clean", "relay_mission_id"])), 6)

    q4 = pd.read_csv(RESULTS / "q4_partition_resources.csv")
    fill_sheet(wb["Q4_分区配置"], list(rows_from(q4,
        ["K", "group", "services", "A_uav", "B_uav", "C_uav", "A_battery", "B_battery", "C_battery", "relay_uav", "relay_pack"])), 11)

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
